"""Read the dealer's leads Excel. Column names are matched loosely (Hebrew/English).

When the real file arrives, add its exact headers to COLUMN_ALIASES if needed.
"""
from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import IO

from openpyxl import load_workbook

from .phone import normalize_il_phone

COLUMN_ALIASES: dict[str, list[str]] = {
    "first_name": ["שם פרטי", "שם", "first name", "firstname", "first", "name"],
    "last_name": ["שם משפחה", "משפחה", "last name", "lastname", "surname", "last"],
    "full_name": ["שם מלא", "full name", "fullname"],
    "email": ["מייל", "אימייל", "דוא\"ל", "דואל", "email", "e-mail", "mail"],
    "phone": ["טלפון", "נייד", "פלאפון", "מספר טלפון", "phone", "mobile", "cell", "telephone"],
    "model": ["מודל", "דגם", "דגם מבוקש", "מודל מבוקש", "model", "interested model", "bike", "product"],
}


def _clean(h: object) -> str:
    return str(h or "").strip().lower().replace("_", " ")


@dataclass
class ParsedLead:
    first_name: str
    last_name: str
    email: str
    phone: str | None  # normalized intl, or None if invalid
    phone_raw: str
    model_raw: str
    row: int


@dataclass
class ImportResult:
    leads: list[ParsedLead] = field(default_factory=list)
    column_map: dict[str, str] = field(default_factory=dict)  # field -> header found
    missing_fields: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)


def map_columns(headers: list[object]) -> dict[str, int]:
    cleaned = [_clean(h) for h in headers]
    result: dict[str, int] = {}
    used: set[int] = set()
    # exact matches first, then "contains" — avoids "שם" grabbing "שם משפחה"
    for exact in (True, False):
        for fld, aliases in COLUMN_ALIASES.items():
            if fld in result:
                continue
            for alias in aliases:
                a = _clean(alias)
                for i, h in enumerate(cleaned):
                    if i in used or not h:
                        continue
                    if (h == a) if exact else (a in h):
                        result[fld] = i
                        used.add(i)
                        break
                if fld in result:
                    break
    return result


def parse_leads_excel(src: str | Path | IO[bytes]) -> ImportResult:
    wb = load_workbook(src, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    res = ImportResult()
    if not rows:
        res.warnings.append("הקובץ ריק.")
        return res

    # header = first row that has at least 2 non-empty cells
    header_idx = next((i for i, r in enumerate(rows) if sum(1 for c in r if c not in (None, "")) >= 2), 0)
    headers = list(rows[header_idx])
    cols = map_columns(headers)
    res.column_map = {f: str(headers[i]) for f, i in cols.items()}

    has_name = "first_name" in cols or "full_name" in cols
    for f in ("phone", "model"):
        if f not in cols:
            res.missing_fields.append(f)
    if not has_name:
        res.missing_fields.append("first_name")

    def cell(r, f):
        i = cols.get(f)
        if i is None or i >= len(r):
            return ""
        v = r[i]
        if v is None:
            return ""
        if isinstance(v, float) and v.is_integer():
            v = int(v)
        return str(v).strip()

    for n, r in enumerate(rows[header_idx + 1 :], start=header_idx + 2):
        if not any(c not in (None, "") for c in r):
            continue
        first, last = cell(r, "first_name"), cell(r, "last_name")
        if not first and cell(r, "full_name"):
            parts = cell(r, "full_name").split()
            first, last = parts[0], " ".join(parts[1:])
        phone_raw = cell(r, "phone")
        phone = normalize_il_phone(phone_raw)
        if phone_raw and not phone:
            res.warnings.append(f"שורה {n}: מספר טלפון לא תקין ({phone_raw}).")
        if not phone_raw:
            res.warnings.append(f"שורה {n}: אין מספר טלפון.")
        res.leads.append(
            ParsedLead(
                first_name=first,
                last_name=last,
                email=cell(r, "email"),
                phone=phone,
                phone_raw=phone_raw,
                model_raw=cell(r, "model"),
                row=n,
            )
        )
    wb.close()
    return res
